import os
import json
import re
import smtplib
import socket
import ssl
import sys
import time
from io import BytesIO
from datetime import date, datetime
from email.message import EmailMessage
from urllib.parse import parse_qs, urlparse
import unicodedata
from flask import Flask, jsonify, redirect, render_template, request, session, flash, url_for
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import UniqueConstraint, text
import requests

app = Flask(__name__)
app.secret_key = os.getenv('SECRET_KEY', 'your-secret-key-change-this')

# Configura o caminho do micro banco (arquivo .db)
basedir = os.path.abspath(os.path.dirname(__file__))
db_path = os.path.join(basedir, 'micro_banco.db')
app.config['SQLALCHEMY_DATABASE_URI'] = os.getenv('DATABASE_URL', f'sqlite:///{db_path}')
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db = SQLAlchemy(app)
produtos_cache = {}

EMPRESAS_DISPONIVEIS = ['Varejo', 'Redemac', 'Alternativa', 'Granteck', 'Especial']

SMTP_HOST = (os.getenv('SMTP_HOST') or 'smtp.gmail.com').strip()
SMTP_PORT = int((os.getenv('SMTP_PORT') or '587').strip())
SMTP_USER = (os.getenv('SMTP_USER') or 'ecalliari@gmail.com').strip()
SMTP_PASSWORD = (os.getenv('SMTP_PASSWORD') or 'ahtl licw snuv tfzw').strip()
SMTP_USE_TLS = (os.getenv('SMTP_USE_TLS') or 'true').strip().lower() in ('1', 'true', 'yes', 'on')



# Modelos do Banco de Dados
class Usuario(db.Model):
    __table_args__ = (
        UniqueConstraint('email', 'cnpj', name='uq_usuario_email_cnpj'),
    )

    id = db.Column(db.Integer, primary_key=True)
    nome = db.Column(db.String(80), nullable=False)
    email = db.Column(db.String(120), nullable=False)
    cnpj = db.Column(db.String(14), nullable=False)
    celular = db.Column(db.String(20), unique=True, nullable=True)
    associacoes = db.Column(db.String(255), nullable=False, default='Varejo')
    empresa_ativa = db.Column(db.String(120), nullable=False, default='Varejo')
    perm_download_tabelas = db.Column(db.Boolean, nullable=False, default=True)
    perm_fazer_pedido = db.Column(db.Boolean, nullable=False, default=True)
    perm_solicitar_visita = db.Column(db.Boolean, nullable=False, default=True)
    perm_visualizar_pedidos = db.Column(db.Boolean, nullable=False, default=True)
    is_admin = db.Column(db.Boolean, nullable=False, default=False)
    senha = db.Column(db.String(120), nullable=True)


class Agendamento(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    usuario_id = db.Column(db.Integer, db.ForeignKey('usuario.id'), nullable=False)
    data_hora = db.Column(db.DateTime, nullable=False)
    observacao = db.Column(db.String(255), nullable=True)


class Pedido(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    usuario_id = db.Column(db.Integer, db.ForeignKey('usuario.id'), nullable=False)
    data_pedido = db.Column(db.DateTime, default=datetime.now, nullable=False)
    valor_total = db.Column(db.Float, nullable=False)
    status = db.Column(db.String(50), default='Pendente', nullable=False)

    itens = db.relationship('ItensPedido', backref='pedido', lazy=True, cascade='all, delete-orphan')


class ItensPedido(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    pedido_id = db.Column(db.Integer, db.ForeignKey('pedido.id'), nullable=False)
    referencia_produto = db.Column(db.String(120), nullable=False)
    quantidade = db.Column(db.Integer, nullable=False)
    desconto = db.Column(db.Float, default=0.0, nullable=True)
    valor_unitario = db.Column(db.Float, nullable=False)
    valor_total = db.Column(db.Float, nullable=False)


class DeletarUsuarios(db.Model):
    __tablename__ = 'deletar_usuarios'
    id = db.Column(db.Integer, primary_key=True)

    @staticmethod
    def deletar_todos():
        num_rows_deleted = db.session.query(Usuario).delete()
        db.session.commit()
        return num_rows_deleted


# Funções Utilitárias e de Normalização
def normalizar_cnpj(cnpj):
    return re.sub(r'\D', '', str(cnpj or ''))


def valor_booleano(valor, padrao=False):
    if valor is None:
        return padrao
    if isinstance(valor, bool):
        return valor
    return str(valor).strip().lower() in ('true', '1', 'on', 'yes')


def parse_empresas(valor_empresas):
    if isinstance(valor_empresas, list):
        candidatos = valor_empresas
    else:
        texto = str(valor_empresas or '')
        candidatos = [item.strip() for item in texto.replace(';', ',').split(',') if item.strip()]

    filtradas = [empresa for empresa in candidatos if empresa in EMPRESAS_DISPONIVEIS]
    if not filtradas:
        return ['Varejo']
    return list(dict.fromkeys(filtradas))


def obter_permissoes_usuario(usuario):
    if usuario.is_admin:
        return {
            'download_tabelas': True,
            'fazer_pedido': True,
            'solicitar_visita': True,
            'visualizar_pedidos': True,
            'atualizar_tabelas': True,
        }

    return {
        'download_tabelas': bool(usuario.perm_download_tabelas),
        'fazer_pedido': bool(usuario.perm_fazer_pedido),
        'solicitar_visita': bool(usuario.perm_solicitar_visita),
        'visualizar_pedidos': bool(usuario.perm_visualizar_pedidos),
        'atualizar_tabelas': False,
    }


def contexto_index():
    user_data = session.get('user_data', {})
    return {
        'user_id': user_data.get('id'),
        'nome': user_data.get('nome', 'Usuário'),
        'permissoes': user_data.get('permissoes', {}),
        'empresas': user_data.get('empresas', []),
        'empresa_ativa': user_data.get('empresa_ativa', ''),
        'is_admin': user_data.get('is_admin', False),
    }


def usuario_tem_permissao(chave):
    if 'user_data' not in session:
        return False

    user_data = session['user_data']
    if user_data.get('is_admin'):
        return True

    permissoes = user_data.get('permissoes', {})
    return bool(permissoes.get(chave, False))


def obter_arquivo_tabela_por_associacao(associacao):
    arquivo_map = {
        'Varejo': 'static/tabelas/tabela_varejo.xlsx',
        'Redemac': 'static/tabelas/tabela_redemac.xlsx',
        'Alternativa': 'static/tabelas/tabela_alternativa.xlsx',
        'Granteck': 'static/tabelas/tabela_granteck.xlsx',
        'Especial': 'static/tabelas/tabela_especial.xlsx'
    }
    return arquivo_map.get(associacao, 'static/tabelas/tabela.xlsx')


def normalizar_texto_coluna(valor):
    texto = str(valor or '').strip().lower()
    texto = unicodedata.normalize('NFKD', texto)
    texto = ''.join(ch for ch in texto if not unicodedata.combining(ch))
    texto = texto.replace('', '')
    texto = re.sub(r'[^a-z0-9]+', '', texto)
    return texto


def encontrar_indice_coluna(cabecalho, aliases):
    aliases_norm = {normalizar_texto_coluna(alias) for alias in aliases}
    for idx, coluna in enumerate(cabecalho):
        if normalizar_texto_coluna(coluna) in aliases_norm:
            return idx
    return None


def normalizar_decimal_generico(valor):
    if valor is None:
        return None

    texto = str(valor).strip()
    if not texto:
        return None

    texto = re.sub(r'[^0-9,.-]', '', texto)
    if ',' in texto and '.' in texto:
        texto = texto.replace('.', '').replace(',', '.')
    elif ',' in texto:
        texto = texto.replace(',', '.')

    try:
        return float(texto)
    except ValueError:
        return None


def normalizar_data_generica(valor):
    if valor is None:
        return None

    if isinstance(valor, datetime):
        return valor.date()

    if isinstance(valor, date):
        return valor

    if isinstance(valor, (int, float)):
        try:
            from openpyxl.utils.datetime import from_excel
            convertido = from_excel(valor)
            if isinstance(convertido, datetime):
                return convertido.date()
            if isinstance(convertido, date):
                return convertido
        except Exception:
            return None

    texto = str(valor).strip()
    if not texto:
        return None

    texto_sem_hora = re.split(r'\s+', texto, maxsplit=1)[0]
    formatos = (
        '%d/%m/%Y',
        '%d/%m/%y',
        '%Y-%m-%d',
        '%d-%m-%Y',
        '%Y/%m/%d',
    )

    for formato in formatos:
        try:
            return datetime.strptime(texto_sem_hora, formato).date()
        except ValueError:
            continue

    return None


def geocodificar_endereco(endereco):
    if not endereco:
        return None, None

    headers = {
        'User-Agent': 'app-vendas/1.0 (sales-map geocoding)'
    }
    params = {
        'q': endereco,
        'format': 'json',
        'limit': 1,
        'countrycodes': 'br',
    }

    try:
        resposta = requests.get(
            'https://nominatim.openstreetmap.org/search',
            params=params,
            headers=headers,
            timeout=15,
        )
    except requests.RequestException:
        return None, None

    if resposta.status_code != 200:
        return None, None

    dados = resposta.json() if resposta.headers.get('Content-Type', '').startswith('application/json') else []
    if not dados:
        return None, None

    lat = normalizar_decimal_generico(dados[0].get('lat'))
    lon = normalizar_decimal_generico(dados[0].get('lon'))
    return lat, lon


def montar_endereco_cliente(endereco, bairro, cidade, estado, cep):
    partes = [
        str(endereco or '').strip(),
        str(bairro or '').strip(),
        str(cidade or '').strip(),
        str(estado or '').strip(),
        str(cep or '').strip(),
        'Brasil',
    ]
    partes_validas = [parte for parte in partes if parte]
    return ', '.join(partes_validas)


def coluna_existe(nome_tabela, nome_coluna):
    try:
        resultado = db.session.execute(text(f'PRAGMA table_info({nome_tabela})'))
        colunas = [linha[1] for linha in resultado.fetchall()]
        return nome_coluna in colunas
    except Exception:
        return False


def atualizar_schema_legado():
    alteracoes = []

    if not coluna_existe('usuario', 'cnpj'):
        alteracoes.append('ALTER TABLE usuario ADD COLUMN cnpj VARCHAR(14)')
    if not coluna_existe('usuario', 'associacoes'):
        alteracoes.append("ALTER TABLE usuario ADD COLUMN associacoes VARCHAR(255) DEFAULT 'Varejo'")
    if not coluna_existe('usuario', 'empresa_ativa'):
        alteracoes.append("ALTER TABLE usuario ADD COLUMN empresa_ativa VARCHAR(120) DEFAULT 'Varejo'")
    if not coluna_existe('usuario', 'perm_download_tabelas'):
        alteracoes.append('ALTER TABLE usuario ADD COLUMN perm_download_tabelas BOOLEAN DEFAULT 1')
    if not coluna_existe('usuario', 'perm_fazer_pedido'):
        alteracoes.append('ALTER TABLE usuario ADD COLUMN perm_fazer_pedido BOOLEAN DEFAULT 1')
    if not coluna_existe('usuario', 'perm_solicitar_visita'):
        alteracoes.append('ALTER TABLE usuario ADD COLUMN perm_solicitar_visita BOOLEAN DEFAULT 1')
    if not coluna_existe('usuario', 'perm_visualizar_pedidos'):
        alteracoes.append('ALTER TABLE usuario ADD COLUMN perm_visualizar_pedidos BOOLEAN DEFAULT 1')
    if not coluna_existe('usuario', 'is_admin'):
        alteracoes.append('ALTER TABLE usuario ADD COLUMN is_admin BOOLEAN DEFAULT 0')
    if not coluna_existe('itens_pedido', 'desconto'):
        alteracoes.append('ALTER TABLE itens_pedido ADD COLUMN desconto FLOAT DEFAULT 0.0')

    for sql in alteracoes:
        try:
            db.session.execute(text(sql))
        except Exception:
            pass

    if alteracoes:
        db.session.commit()

    if coluna_existe('usuario', 'associacao'):
        db.session.execute(text("""
            UPDATE usuario
            SET associacoes = COALESCE(NULLIF(TRIM(associacao), ''), 'Varejo')
            WHERE associacoes IS NULL OR TRIM(associacoes) = ''
        """))
        db.session.execute(text("""
            UPDATE usuario
            SET empresa_ativa = COALESCE(NULLIF(TRIM(associacao), ''), 'Varejo')
            WHERE empresa_ativa IS NULL OR TRIM(empresa_ativa) = ''
        """))

    db.session.execute(text("""
        UPDATE usuario
        SET cnpj = '00000000000000'
        WHERE cnpj IS NULL OR TRIM(cnpj) = ''
    """))

    db.session.commit()


def garantir_usuario_admin():
    admin_email = os.getenv('ADMIN_EMAIL', 'admin@sistema.local').strip().lower()
    admin_cnpj = normalizar_cnpj(os.getenv('ADMIN_CNPJ', '00000000000000'))
    admin_senha = os.getenv('ADMIN_SENHA', 'admin123')
    empresas = ','.join(EMPRESAS_DISPONIVEIS)

    admin = Usuario.query.filter_by(email=admin_email, cnpj=admin_cnpj).first()
    if admin is None:
        admin = Usuario(
            nome='Administrador',
            email=admin_email,
            cnpj=admin_cnpj,
            celular='(00)00000-0000',
            associacoes=empresas,
            empresa_ativa=EMPRESAS_DISPONIVEIS[0],
            senha=admin_senha,
            perm_download_tabelas=True,
            perm_fazer_pedido=True,
            perm_solicitar_visita=True,
            perm_visualizar_pedidos=True,
            is_admin=True,
        )
        db.session.add(admin)
    else:
        admin.is_admin = True
        admin.perm_download_tabelas = True
        admin.perm_fazer_pedido = True
        admin.perm_solicitar_visita = True
        admin.perm_visualizar_pedidos = True
        admin.associacoes = empresas
        admin.empresa_ativa = EMPRESAS_DISPONIVEIS[0]
        if not admin.senha:
            admin.senha = admin_senha

    db.session.commit()


# Inicialização das tabelas e garantia de pastas estáticas
with app.app_context():
    os.makedirs(os.path.join(basedir, 'static', 'tabelas'), exist_ok=True)
    os.makedirs(os.path.join(basedir, 'static', 'vendas'), exist_ok=True)
    os.makedirs(os.path.join(basedir, 'static', 'carteira'), exist_ok=True)
    db.create_all()
    atualizar_schema_legado()
    garantir_usuario_admin()


def serializar_usuario(usuario):
    empresas = parse_empresas(usuario.associacoes)
    empresa_ativa = usuario.empresa_ativa if usuario.empresa_ativa in empresas else empresas[0]

    return {
        'id': usuario.id,
        'nome': usuario.nome,
        'email': usuario.email,
        'cnpj': usuario.cnpj,
        'celular': usuario.celular,
        'empresas': empresas,
        'empresa_ativa': empresa_ativa,
        'permissoes': obter_permissoes_usuario(usuario),
        'is_admin': bool(usuario.is_admin),
    }


def serializar_pedido(pedido):
    usuario = Usuario.query.filter_by(id=pedido.usuario_id).first()
    return {
        'id': pedido.id,
        'usuario_id': pedido.usuario_id,
        'cliente_nome': usuario.nome if usuario else '',
        'cliente_email': usuario.email if usuario else '',
        'cliente_cnpj': usuario.cnpj if usuario else '',
        'data_pedido': pedido.data_pedido.isoformat(),
        'valor_total': round(pedido.valor_total, 2),
        'status': pedido.status,
        'quantidade_itens': len(pedido.itens),
    }


def usuario_autenticado_sessao():
    user_data = session.get('user_data') or {}
    user_id = user_data.get('id')
    if not user_id:
        return None
    return Usuario.query.filter_by(id=user_id).first()


def usuario_solicitante_da_requisicao(payload=None):
    if payload is None:
        payload = request.get_json(silent=True)

    if not payload:
        payload = request.args or request.form

    usuario_id = None
    if hasattr(payload, 'get'):
        usuario_id = payload.get('usuario_id') or payload.get('solicitante_id')

    if usuario_id:
        try:
            return Usuario.query.filter_by(id=int(usuario_id)).first()
        except (TypeError, ValueError):
            return None

    return usuario_autenticado_sessao()


def pode_acessar_pedido(usuario_solicitante, pedido):
    return bool(usuario_solicitante and (usuario_solicitante.is_admin or pedido.usuario_id == usuario_solicitante.id))


def obter_ips_ipv4(host, port):
    try:
        infos = socket.getaddrinfo(host, port, socket.AF_INET, socket.SOCK_STREAM)
        ips = [info[4][0] for info in infos if info and len(info) > 4]
        return list(dict.fromkeys(ips))
    except Exception:
        return [host]


def conectar_smtp_resiliente(host, port, user, password, use_tls=True):
    senha_limpa = (password or '').replace(' ', '').strip()
    erros = []

    # 1. Prioridade absoluta: Porta 465 (SSL) em IPv4 com check_hostname = False
    # Isso evita bloqueios de porta 587 no Render e erros de IPv6 / certificado IP mismatch.
    ips_465 = obter_ips_ipv4(host, 465)
    for ip in ips_465:
        try:
            context = ssl.create_default_context()
            context.check_hostname = False
            servidor = smtplib.SMTP_SSL(ip, 465, timeout=15, context=context)
            if user and senha_limpa:
                servidor.login(user, senha_limpa)
            return servidor
        except Exception as e:
            erros.append(f"SSL IPv4 {ip}:465 -> {e}")

    # 2. Tentar porta 587 (TLS com IPv4) se porta 465 falhar
    ips_587 = obter_ips_ipv4(host, 587)
    for ip in ips_587:
        try:
            context = ssl.create_default_context()
            context.check_hostname = False
            servidor = smtplib.SMTP(ip, 587, timeout=15)
            if use_tls:
                servidor.starttls(context=context)
            if user and senha_limpa:
                servidor.login(user, senha_limpa)
            return servidor
        except Exception as e:
            erros.append(f"TLS IPv4 {ip}:587 -> {e}")

    # 3. Fallback final usando o hostname direto
    try:
        context = ssl.create_default_context()
        context.check_hostname = False
        servidor = smtplib.SMTP_SSL(host, 465, timeout=15, context=context)
        if user and senha_limpa:
            servidor.login(user, senha_limpa)
        return servidor
    except Exception as e:
        erros.append(f"Direct {host}:465 -> {e}")

    raise RuntimeError(f"Falha ao conectar no servidor SMTP: {'; '.join(erros)}")


def enviar_email_pedido_admin(usuario, pedido, itens):
    recipientes = []
    if usuario and usuario.email and '@' in usuario.email and not usuario.email.endswith('.local'):
        recipientes.append(usuario.email)

    if SMTP_USER and SMTP_USER not in recipientes and '@' in SMTP_USER:
        recipientes.append(SMTP_USER)

    if not recipientes:
        if usuario and usuario.email:
            recipientes.append(usuario.email)
        elif SMTP_USER:
            recipientes.append(SMTP_USER)
        else:
            raise RuntimeError('Usuário sem email válido para receber confirmação do pedido.')

    assunto = f'Confirmacao do pedido #{pedido.id} - {usuario.nome if usuario else "Cliente"}'
    linhas_itens = []
    for item in itens:
        desconto = float(item.get('desconto', 0.0) or 0.0)
        vlr_unit = float(item.get('valor_unitario', 0))
        vlr_total = float(item.get('valor_total', 0))
        str_desc = f" | Desc: {desconto:.1f}%" if desconto > 0 else ""
        linhas_itens.append(
            f"- Ref: {item.get('referencia_produto')} | Qtd: {item.get('quantidade')}{str_desc} | "
            f"Vlr Unit.: R$ {vlr_unit:.2f} | "
            f"Vlr Total: R$ {vlr_total:.2f}"
        )

    corpo = (
        f"Seu pedido foi finalizado com sucesso na plataforma.\n\n"
        f"Pedido: #{pedido.id}\n"
        f"Data: {pedido.data_pedido.strftime('%d/%m/%Y %H:%M:%S')}\n"
        f"Valor total: R$ {pedido.valor_total:.2f}\n\n"
        f"Dados do cliente:\n"
        f"Nome: {usuario.nome if usuario else '-'}\n"
        f"Email: {usuario.email if usuario else '-'}\n"
        f"CNPJ: {usuario.cnpj if usuario else '-'}\n"
        f"Celular: {(usuario.celular if usuario else None) or '-'}\n"
        f"Empresa ativa: {usuario.empresa_ativa if usuario else '-'}\n"
        f"Empresas habilitadas: {usuario.associacoes if usuario else '-'}\n\n"
        f"Itens do pedido:\n"
        f"{'\n'.join(linhas_itens)}\n"
    )

    mensagem = EmailMessage()
    mensagem['Subject'] = assunto
    mensagem['From'] = SMTP_USER or 'no-reply@sistema.local'
    mensagem['To'] = ', '.join(recipientes)
    mensagem.set_content(corpo)

    if not SMTP_HOST:
        raise RuntimeError('SMTP_HOST não configurado para envio de email.')

    servidor = conectar_smtp_resiliente(
        host=SMTP_HOST,
        port=SMTP_PORT,
        user=SMTP_USER,
        password=SMTP_PASSWORD,
        use_tls=SMTP_USE_TLS
    )
    try:
        servidor.send_message(mensagem)
    finally:
        try:
            servidor.quit()
        except Exception:
            pass


# Rotas Principais do Frontend (Interface Web)
@app.route('/')
def home():
    if 'user_data' in session:
        return render_template('index.html', **contexto_index())
    return render_template('home.html')


@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'GET':
        return render_template('login_cadastro.html')

    is_json = request.is_json or request.headers.get('Content-Type', '').startswith('application/json')
    payload = request.get_json(silent=True) or request.form

    email = (payload.get('email') or '').strip().lower()
    cnpj = normalizar_cnpj(payload.get('cnpj'))
    senha = payload.get('senha')

    if not all([email, senha]):
        erro = 'Email e senha são obrigatórios.'
        if is_json:
            return jsonify({'erro': erro}), 400
        flash(erro, 'error')
        return render_template('login_cadastro.html', error=erro)

    usuario = None
    if cnpj:
        usuario = Usuario.query.filter_by(email=email, cnpj=cnpj).first()

    if usuario is None:
        usuario = Usuario.query.filter_by(email=email).first()

    if usuario is None:
        erro = 'Usuário não encontrado para este email.'
        if is_json:
            return jsonify({'erro': erro}), 404
        flash(erro, 'error')
        return render_template('login_cadastro.html', error=erro)

    if usuario.senha != senha:
        erro = 'Senha incorreta.'
        if is_json:
            return jsonify({'erro': erro}), 401
        flash(erro, 'error')
        return render_template('login_cadastro.html', error=erro)

    session['user_data'] = serializar_usuario(usuario)

    if is_json:
        return jsonify(session['user_data']), 200

    return render_template('index.html', **contexto_index())


@app.route('/adicionar', methods=['GET', 'POST'])
def adicionar():
    if request.method == 'GET':
        return render_template('login_cadastro.html')

    is_json = request.is_json or request.headers.get('Content-Type', '').startswith('application/json')
    payload = request.get_json(silent=True) or request.form

    nome = (payload.get('nome') or '').strip()
    email = (payload.get('email') or '').strip().lower()
    celular = (payload.get('celular') or '').strip()
    cnpj = normalizar_cnpj(payload.get('cnpj'))
    senha = payload.get('senha')

    empresas = payload.getlist('empresas') if hasattr(payload, 'getlist') and payload.getlist('empresas') else payload.get('empresas', 'Varejo')
    empresas_lista = parse_empresas(empresas)
    empresas_salvar = ','.join(empresas_lista)
    empresa_ativa = empresas_lista[0]

    perm_download = 'perm_download_tabelas' in request.form if not is_json else valor_booleano(payload.get('perm_download_tabelas', True), True)
    perm_pedido = 'perm_fazer_pedido' in request.form if not is_json else valor_booleano(payload.get('perm_fazer_pedido', True), True)
    perm_visita = 'perm_solicitar_visita' in request.form if not is_json else valor_booleano(payload.get('perm_solicitar_visita', True), True)
    perm_pedidos = 'perm_visualizar_pedidos' in request.form if not is_json else valor_booleano(payload.get('perm_visualizar_pedidos', True), True)

    if not all([nome, email, celular, cnpj, senha]):
        erro = 'Todos os campos são obrigatórios.'
        if is_json:
            return jsonify({'erro': erro}), 400
        flash(erro, 'error')
        return render_template('login_cadastro.html', error=erro)

    if len(cnpj) != 14:
        erro = 'CNPJ inválido. Informe 14 dígitos.'
        if is_json:
            return jsonify({'erro': erro}), 400
        flash(erro, 'error')
        return render_template('login_cadastro.html', error=erro)

    if Usuario.query.filter_by(email=email, cnpj=cnpj).first():
        erro = 'Já existe um usuário cadastrado com este email e CNPJ.'
        if is_json:
            return jsonify({'erro': erro}), 409
        flash(erro, 'error')
        return render_template('login_cadastro.html', error=erro)

    if Usuario.query.filter(Usuario.celular == celular).first():
        erro = 'Já existe um usuário cadastrado com este celular.'
        if is_json:
            return jsonify({'erro': erro}), 409
        flash(erro, 'error')
        return render_template('login_cadastro.html', error=erro)

    novo_usuario = Usuario(
        nome=nome,
        email=email,
        cnpj=cnpj,
        celular=celular,
        associacoes=empresas_salvar,
        empresa_ativa=empresa_ativa,
        perm_download_tabelas=perm_download,
        perm_fazer_pedido=perm_pedido,
        perm_solicitar_visita=perm_visita,
        perm_visualizar_pedidos=perm_pedidos,
        is_admin=False,
        senha=senha,
    )
    db.session.add(novo_usuario)
    db.session.commit()

    usuario_criado = serializar_usuario(novo_usuario)
    session['user_data'] = usuario_criado

    if is_json:
        return jsonify(usuario_criado), 201

    return render_template('index.html', **contexto_index())


@app.route('/logout')
def logout():
    session.pop('user_data', None)
    return redirect(url_for('home'))


@app.route('/agendar', methods=['GET', 'POST'])
def agendar():
    if 'user_data' not in session:
        flash('Por favor, faça login para acessar a página de agendamento.', 'error')
        return render_template('login_cadastro.html', error='Faça login para acessar o agendamento.')

    if not usuario_tem_permissao('solicitar_visita'):
        flash('Você não possui permissão para solicitar visita.', 'error')
        return render_template('index.html', **contexto_index())

    if request.method == 'POST':
        usuario_id = session['user_data']['id']
        data_hora_raw = request.form.get('data')
        observacao = request.form.get('observacoes', '')

        try:
            data_hora = datetime.fromisoformat(str(data_hora_raw))
        except (ValueError, TypeError):
            flash('Formato de data inválido.', 'error')
            return render_template('index.html', error='Data inválida', **contexto_index())

        novo_agendamento = Agendamento(usuario_id=usuario_id, data_hora=data_hora, observacao=observacao)
        db.session.add(novo_agendamento)
        db.session.commit()

        flash('Agendamento realizado com sucesso!', 'success')

    return render_template('index.html', **contexto_index())


@app.route('/adicionaragenda', methods=['POST'])
def adicionaragenda():
    payload = request.get_json(silent=True) or request.form
    usuario_id = payload.get('usuario_id')
    data_hora_raw = payload.get('data_hora')
    observacao = payload.get('observacao', '')

    if not usuario_id or not data_hora_raw:
        return jsonify({'erro': 'usuario_id e data_hora são obrigatórios.'}), 400

    try:
        data_hora = datetime.fromisoformat(str(data_hora_raw))
    except ValueError:
        return jsonify({'erro': 'Formato de data inválido. Use YYYY-MM-DD ou ISO completo.'}), 400

    novo_agendamento = Agendamento(usuario_id=usuario_id, data_hora=data_hora, observacao=observacao)
    db.session.add(novo_agendamento)
    db.session.commit()

    return jsonify({'mensagem': 'Agendamento realizado com sucesso!'}), 201


@app.route('/carregar_produtos', methods=['GET'])
def carregar_produtos():
    try:
        if 'user_data' not in session and not request.args.get('usuario_id'):
            return jsonify({'erro': 'Usuário não autenticado'}), 401

        from openpyxl import load_workbook

        associacao = request.args.get('associacao') or request.args.get('empresa')
        usuario_id = request.args.get('usuario_id', type=int)

        if usuario_id:
            usuario = Usuario.query.filter_by(id=usuario_id).first()
            if not usuario:
                return jsonify({'erro': 'Usuário não encontrado'}), 404
            if not usuario.is_admin and not usuario.perm_download_tabelas:
                return jsonify({'erro': 'Usuário sem permissão para baixar tabelas'}), 403
            empresas_usuario = parse_empresas(usuario.associacoes)
            if associacao and associacao not in empresas_usuario:
                return jsonify({'erro': 'Empresa não habilitada para este usuário'}), 403
            if not associacao:
                associacao = usuario.empresa_ativa
        elif 'user_data' in session:
            usuario_info = session['user_data']
            if not usuario_info.get('is_admin') and not usuario_tem_permissao('download_tabelas'):
                return jsonify({'erro': 'Sem permissão para baixar tabelas'}), 403
            if not associacao:
                associacao = usuario_info.get('empresa_ativa', 'Varejo')

        if not associacao:
            associacao = 'Varejo'

        arquivo = obter_arquivo_tabela_por_associacao(associacao)
        arquivo_caminho = os.path.join(basedir, arquivo)
        if not os.path.exists(arquivo_caminho):
            return jsonify({'erro': f'Arquivo {arquivo} não encontrado'}), 404

        mtime = os.path.getmtime(arquivo_caminho)
        cache_atual = produtos_cache.get(associacao)
        if cache_atual and cache_atual.get('mtime') == mtime:
            return jsonify(cache_atual.get('produtos', [])), 200

        workbook = load_workbook(filename=arquivo_caminho, data_only=True, read_only=True)
        sheet = workbook.active

        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return jsonify({'erro': 'Arquivo de produtos vazio'}), 400

        cabecalho = [str(c or '').strip() for c in rows[0]]

        def indice_coluna(nome):
            nome_normalizado = nome.strip().lower()
            for idx, col in enumerate(cabecalho):
                if col.strip().lower() == nome_normalizado:
                    return idx
            return None

        idx_referencia = indice_coluna('Referência')
        idx_descricao = indice_coluna('Descrição')
        idx_emb = indice_coluna('Emb.')
        idx_preco = indice_coluna('Preço final')

        if None in [idx_referencia, idx_descricao, idx_emb, idx_preco]:
            return jsonify({'erro': 'Formato de arquivo inválido. Necessário: Referência, Descrição, Emb., Preço final'}), 400

        produtos = []
        for row in rows[1:]:
            referencia = row[idx_referencia] if idx_referencia < len(row) else None
            descricao = row[idx_descricao] if idx_descricao < len(row) else None
            quantidade_emb = row[idx_emb] if idx_emb < len(row) else None
            preco = row[idx_preco] if idx_preco < len(row) else None

            if referencia is None or quantidade_emb is None or preco is None:
                continue

            valor = normalizar_decimal_generico(preco)
            quantidade_embalagem = normalizar_decimal_generico(quantidade_emb)

            if valor is None or quantidade_embalagem is None:
                continue

            produtos.append({
                'referencia': str(referencia).strip(),
                'descricao': str(descricao).strip() if descricao is not None else '',
                'quantidade_embalagem': round(quantidade_embalagem, 2),
                'valor': round(valor, 2)
            })

        produtos_cache[associacao] = {
            'mtime': mtime,
            'produtos': produtos
        }

        return jsonify(produtos), 200
    except Exception as e:
        return jsonify({'erro': str(e)}), 500


@app.route('/atualizar_tabelas', methods=['POST'])
def atualizar_tabelas():
    if 'user_data' not in session and not request.form.get('usuario_id'):
        flash('Usuário não autenticado.', 'error')
        return redirect('/')

    user_data = session.get('user_data', {})
    is_admin = user_data.get('is_admin', False)

    usuario_id = request.form.get('usuario_id', type=int) or user_data.get('id')
    associacao = (request.form.get('associacao') or request.form.get('empresa') or request.form.get('empresa_arquivo') or '').strip()

    if usuario_id and not is_admin:
        usuario = Usuario.query.filter_by(id=usuario_id).first()
        if usuario:
            is_admin = usuario.is_admin

    if not is_admin:
        if request.is_json:
            return jsonify({'erro': 'Apenas administrador pode atualizar tabelas.'}), 403
        flash('Apenas administrador pode atualizar tabelas.', 'error')
        return render_template('index.html', **contexto_index())

    arquivo_upload = request.files.get('arquivo_tabela')
    if not associacao:
        if request.is_json:
            return jsonify({'erro': 'Empresa é obrigatória.'}), 400
        flash('Selecione a empresa para atualização.', 'error')
        return render_template('index.html', **contexto_index())

    if arquivo_upload is None or not arquivo_upload.filename:
        if request.is_json:
            return jsonify({'erro': 'Arquivo .xlsx é obrigatório.'}), 400
        flash('Selecione um arquivo .xlsx para atualizar.', 'error')
        return render_template('index.html', **contexto_index())

    try:
        nome_arquivo = arquivo_upload.filename.lower().strip()
        if not nome_arquivo.endswith('.xlsx'):
            if request.is_json:
                return jsonify({'erro': 'Formato inválido. Envie um arquivo .xlsx'}), 400
            flash('Formato inválido. Envie um arquivo .xlsx', 'error')
            return render_template('index.html', **contexto_index())

        arquivo_destino_rel = obter_arquivo_tabela_por_associacao(associacao)
        arquivo_destino = os.path.join(basedir, arquivo_destino_rel)

        os.makedirs(os.path.dirname(arquivo_destino), exist_ok=True)

        if os.path.exists(arquivo_destino):
            pasta_destino = os.path.dirname(arquivo_destino)
            base_nome, extensao = os.path.splitext(os.path.basename(arquivo_destino))
            data_hoje = datetime.now().strftime('%Y%m%d')

            backup_nome = f'{base_nome}_{data_hoje}{extensao}'
            backup_path = os.path.join(pasta_destino, backup_nome)
            contador = 1

            while os.path.exists(backup_path):
                backup_nome = f'{base_nome}_{data_hoje}_{contador}{extensao}'
                backup_path = os.path.join(pasta_destino, backup_nome)
                contador += 1

            try:
                os.rename(arquivo_destino, backup_path)
            except Exception:
                pass

        arquivo_upload.save(arquivo_destino)
        produtos_cache.pop(associacao, None)

        msg = f'Tabela da empresa {associacao} atualizada com sucesso!'
        if request.is_json:
            return jsonify({'mensagem': msg, 'arquivo': arquivo_destino_rel}), 200

        flash(msg, 'success')
        return render_template('index.html', **contexto_index())
    except Exception as e:
        if request.is_json:
            return jsonify({'erro': str(e)}), 500
        flash(f'Erro ao atualizar tabela: {str(e)}', 'error')
        return render_template('index.html', **contexto_index())


@app.route('/dados_mapa_vendas', methods=['GET'])
def dados_mapa_vendas():
    if 'user_data' not in session and not request.args.get('usuario_id'):
        return jsonify({'erro': 'Usuário não autenticado'}), 401

    usuario_id = request.args.get('usuario_id', type=int) or session.get('user_data', {}).get('id')
    usuario = Usuario.query.filter_by(id=usuario_id).first() if usuario_id else None

    if not usuario or not usuario.is_admin:
        return jsonify({'erro': 'Apenas administrador pode acessar o mapa de vendas'}), 403

    try:
        from openpyxl import load_workbook

        data_inicio_raw = (request.args.get('data_inicio') or '').strip()
        data_fim_raw = (request.args.get('data_fim') or '').strip()

        data_inicio = None
        data_fim = None
        if data_inicio_raw:
            try:
                data_inicio = datetime.strptime(data_inicio_raw, '%Y-%m-%d').date()
            except ValueError:
                return jsonify({'erro': 'data_inicio inválida. Use o formato YYYY-MM-DD.'}), 400

        if data_fim_raw:
            try:
                data_fim = datetime.strptime(data_fim_raw, '%Y-%m-%d').date()
            except ValueError:
                return jsonify({'erro': 'data_fim inválida. Use o formato YYYY-MM-DD.'}), 400

        if data_inicio and data_fim and data_inicio > data_fim:
            return jsonify({'erro': 'data_inicio não pode ser maior que data_fim.'}), 400

        arquivo_vendas = os.path.join(basedir, 'static', 'vendas', 'vendas_feitas.xlsx')
        arquivo_carteira = os.path.join(basedir, 'static', 'carteira', 'carteira_clientes.xlsx')

        if not os.path.exists(arquivo_vendas):
            return jsonify({'erro': 'Arquivo static/vendas/vendas_feitas.xlsx não encontrado.'}), 404
        if not os.path.exists(arquivo_carteira):
            return jsonify({'erro': 'Arquivo static/carteira/carteira_clientes.xlsx não encontrado.'}), 404

        wb_vendas = load_workbook(filename=arquivo_vendas, data_only=True, read_only=True)
        ws_vendas = wb_vendas.active
        linhas_vendas = list(ws_vendas.iter_rows(values_only=True))
        if not linhas_vendas:
            return jsonify({'erro': 'Planilha de vendas está vazia.'}), 400

        cabecalho_vendas = [str(coluna or '').strip() for coluna in linhas_vendas[0]]
        idx_codigo_venda = encontrar_indice_coluna(cabecalho_vendas, ['Codigo', 'Código'])
        idx_nome_venda = encontrar_indice_coluna(cabecalho_vendas, ['Nome Cliente'])
        idx_valor_venda = encontrar_indice_coluna(cabecalho_vendas, ['Valor Produtos'])
        idx_data_emissao = encontrar_indice_coluna(cabecalho_vendas, ['Data Emissao', 'Data Emissão'])

        if idx_nome_venda is None or idx_valor_venda is None:
            return jsonify({'erro': 'Colunas obrigatórias em vendas_feitas.xlsx: Código, Nome Cliente, Valor Produtos.'}), 400

        if (data_inicio or data_fim) and idx_data_emissao is None:
            return jsonify({'erro': 'Para filtrar por período, a coluna Data Emissao é obrigatória em vendas_feitas.xlsx.'}), 400

        vendas_por_cliente = {}
        for linha in linhas_vendas[1:]:
            nome_cliente = str(linha[idx_nome_venda] or '').strip() if idx_nome_venda < len(linha) else ''
            if not nome_cliente:
                continue

            if data_inicio or data_fim:
                valor_data_emissao = linha[idx_data_emissao] if idx_data_emissao is not None and idx_data_emissao < len(linha) else None
                data_emissao = normalizar_data_generica(valor_data_emissao)
                if data_emissao is None:
                    continue
                if data_inicio and data_emissao < data_inicio:
                    continue
                if data_fim and data_emissao > data_fim:
                    continue

            codigo = str(linha[idx_codigo_venda] or '').strip() if idx_codigo_venda is not None and idx_codigo_venda < len(linha) else ''
            valor_bruto = linha[idx_valor_venda] if idx_valor_venda < len(linha) else None
            valor = normalizar_decimal_generico(valor_bruto) or 0.0

            chave = normalizar_texto_coluna(nome_cliente)
            if chave not in vendas_por_cliente:
                vendas_por_cliente[chave] = {
                    'codigo': codigo,
                    'nome_cliente': nome_cliente,
                    'valor_produtos': 0.0,
                }
            vendas_por_cliente[chave]['valor_produtos'] += valor

        wb_carteira = load_workbook(filename=arquivo_carteira, data_only=True, read_only=True)
        ws_carteira = wb_carteira.active
        linhas_carteira = list(ws_carteira.iter_rows(values_only=True))
        if not linhas_carteira:
            return jsonify({'erro': 'Planilha de carteira está vazia.'}), 400

        cabecalho_carteira = [str(coluna or '').strip() for coluna in linhas_carteira[0]]
        idx_codigo_cli = encontrar_indice_coluna(cabecalho_carteira, ['Codigo', 'Código'])
        idx_nome_cli = encontrar_indice_coluna(cabecalho_carteira, ['Nome Cliente'])
        idx_endereco = encontrar_indice_coluna(cabecalho_carteira, ['Endereco', 'Endereço'])
        idx_bairro = encontrar_indice_coluna(cabecalho_carteira, ['Bairro'])
        idx_cidade = encontrar_indice_coluna(cabecalho_carteira, ['Cidade'])
        idx_estado = encontrar_indice_coluna(cabecalho_carteira, ['Estado', 'UF'])
        idx_cep = encontrar_indice_coluna(cabecalho_carteira, ['CEP'])
        idx_lat = encontrar_indice_coluna(cabecalho_carteira, ['Latitude', 'Lat'])
        idx_lon = encontrar_indice_coluna(cabecalho_carteira, ['Longitude', 'Lon'])

        if idx_nome_cli is None:
            return jsonify({'erro': 'Coluna Nome Cliente não encontrada em carteira_clientes.xlsx.'}), 400

        usar_geocodificacao_fallback = idx_lat is None or idx_lon is None
        cache_geo_path = os.path.join(basedir, 'static', 'carteira', 'geocode_cache.json')
        geocode_cache = {}
        if usar_geocodificacao_fallback and os.path.exists(cache_geo_path):
            try:
                with open(cache_geo_path, 'r', encoding='utf-8') as arquivo_cache:
                    geocode_cache = json.load(arquivo_cache)
            except Exception:
                geocode_cache = {}

        cache_alterado = False
        pontos = []
        clientes_sem_coordenada = 0

        for linha in linhas_carteira[1:]:
            nome_cliente = str(linha[idx_nome_cli] or '').strip() if idx_nome_cli < len(linha) else ''
            if not nome_cliente:
                continue

            codigo = str(linha[idx_codigo_cli] or '').strip() if idx_codigo_cli is not None and idx_codigo_cli < len(linha) else ''
            endereco = linha[idx_endereco] if idx_endereco is not None and idx_endereco < len(linha) else ''
            bairro = linha[idx_bairro] if idx_bairro is not None and idx_bairro < len(linha) else ''
            cidade = linha[idx_cidade] if idx_cidade is not None and idx_cidade < len(linha) else ''
            estado = linha[idx_estado] if idx_estado is not None and idx_estado < len(linha) else ''
            cep = linha[idx_cep] if idx_cep is not None and idx_cep < len(linha) else ''
            endereco_completo = montar_endereco_cliente(endereco, bairro, cidade, estado, cep)

            lat = normalizar_decimal_generico(linha[idx_lat]) if idx_lat is not None and idx_lat < len(linha) else None
            lon = normalizar_decimal_generico(linha[idx_lon]) if idx_lon is not None and idx_lon < len(linha) else None

            cache_key = f"{normalizar_texto_coluna(nome_cliente)}|{normalizar_texto_coluna(endereco_completo)}"
            if (lat is None or lon is None) and cache_key in geocode_cache:
                lat = geocode_cache[cache_key].get('lat')
                lon = geocode_cache[cache_key].get('lon')

            if (lat is None or lon is None) and endereco_completo:
                lat_geo, lon_geo = geocodificar_endereco(endereco_completo)
                if lat_geo is not None and lon_geo is not None:
                    lat = lat_geo
                    lon = lon_geo
                    geocode_cache[cache_key] = {'lat': lat_geo, 'lon': lon_geo}
                    cache_alterado = True

            if lat is None or lon is None:
                clientes_sem_coordenada += 1
                continue

            chave_nome = normalizar_texto_coluna(nome_cliente)
            venda_cliente = vendas_por_cliente.get(chave_nome)
            valor_produtos = round(float(venda_cliente['valor_produtos']), 2) if venda_cliente else 0.0
            fez_venda = venda_cliente is not None

            pontos.append({
                'codigo': codigo,
                'nome_cliente': nome_cliente,
                'valor_produtos': valor_produtos,
                'endereco': endereco_completo,
                'latitude': lat,
                'longitude': lon,
                'fez_venda': fez_venda,
                'cor': 'green' if fez_venda else 'red',
            })

        if cache_alterado:
            try:
                with open(cache_geo_path, 'w', encoding='utf-8') as arquivo_cache:
                    json.dump(geocode_cache, arquivo_cache, ensure_ascii=False, indent=2)
            except Exception:
                pass

        pontos_com_venda = sum(1 for p in pontos if p.get('fez_venda'))
        pontos_sem_venda = sum(1 for p in pontos if not p.get('fez_venda'))

        return jsonify({
            'resumo': {
                'total_clientes_carteira': len(linhas_carteira) - 1,
                'clientes_com_venda': pontos_com_venda,
                'clientes_sem_venda': pontos_sem_venda,
                'clientes_sem_coordenada': clientes_sem_coordenada,
                'periodo': {
                    'data_inicio': data_inicio.isoformat() if data_inicio else None,
                    'data_fim': data_fim.isoformat() if data_fim else None,
                },
            },
            'pontos': pontos,
        }), 200
    except Exception as e:
        return jsonify({'erro': str(e)}), 500


@app.route('/salvar_pedido', methods=['POST'])
def salvar_pedido():
    try:
        if 'user_data' not in session and not request.json.get('usuario_id'):
            return jsonify({'erro': 'Usuário não autenticado'}), 401

        dados = request.get_json(silent=True) or {}
        usuario_solicitante = usuario_solicitante_da_requisicao(dados)

        if not usuario_solicitante:
            return jsonify({'erro': 'Usuário não autenticado'}), 401

        if not usuario_solicitante.is_admin and not usuario_solicitante.perm_fazer_pedido:
            return jsonify({'erro': 'Usuário sem permissão para fazer pedidos'}), 403

        cliente_id = dados.get('cliente_id') or usuario_solicitante.id
        usuario = Usuario.query.filter_by(id=cliente_id).first()
        if not usuario:
            return jsonify({'erro': 'Cliente não encontrado'}), 404

        if not usuario_solicitante.is_admin and usuario.id != usuario_solicitante.id:
            return jsonify({'erro': 'Sem permissão para salvar pedido para outro cliente'}), 403

        itens = dados.get('itens', [])
        if not itens:
            return jsonify({'erro': 'Nenhum item no pedido'}), 400

        valor_total = sum(item['valor_total'] for item in itens)

        novo_pedido = Pedido(usuario_id=usuario.id, valor_total=valor_total)
        db.session.add(novo_pedido)
        db.session.flush()

        for item in itens:
            desconto = float(item.get('desconto', 0.0) or 0.0)
            novo_item = ItensPedido(
                pedido_id=novo_pedido.id,
                referencia_produto=item['referencia_produto'],
                quantidade=item['quantidade'],
                desconto=desconto,
                valor_unitario=item['valor_unitario'],
                valor_total=item['valor_total']
            )
            db.session.add(novo_item)

        db.session.commit()

        aviso_email = None
        if usuario:
            try:
                enviar_email_pedido_admin(usuario, novo_pedido, itens)
            except Exception as email_error:
                import sys
                print(f"[ERRO EMAIL] Falha ao enviar email do pedido #{novo_pedido.id}: {email_error}", file=sys.stderr, flush=True)
                aviso_email = f'Pedido salvo, mas não foi possível enviar email: {str(email_error)}'

        return jsonify({
            'mensagem': 'Pedido salvo com sucesso!',
            'pedido_id': novo_pedido.id,
            'cliente_id': usuario.id,
            'cliente_nome': usuario.nome,
            'valor_total': round(valor_total, 2),
            'aviso_email': aviso_email,
        }), 201
    except Exception as e:
        db.session.rollback()
        return jsonify({'erro': str(e)}), 500


@app.route('/meus_pedidos', methods=['GET'])
@app.route('/pedidos_usuario/<int:usuario_id>', methods=['GET'])
def listar_meus_pedidos(usuario_id=None):
    try:
        if 'user_data' not in session and not request.args.get('usuario_id'):
            return jsonify({'erro': 'Usuário não autenticado'}), 401

        usuario_solicitante = usuario_solicitante_da_requisicao(request.args)
        if not usuario_solicitante:
            return jsonify({'erro': 'Usuário não autenticado'}), 401

        target_id = usuario_id or session.get('user_data', {}).get('id')
        if usuario_id and not usuario_solicitante.is_admin and usuario_solicitante.id != usuario_id:
            return jsonify({'erro': 'Sem permissão para visualizar estes pedidos'}), 403

        pedidos = (
            Pedido.query
            .filter_by(usuario_id=target_id)
            .order_by(Pedido.data_pedido.desc())
            .all()
        )

        retorno = []
        for pedido in pedidos:
            retorno.append({
                'id': pedido.id,
                'data_pedido': pedido.data_pedido.isoformat(),
                'valor_total': round(pedido.valor_total, 2),
                'status': pedido.status,
                'quantidade_itens': len(pedido.itens)
            })

        return jsonify(retorno), 200
    except Exception as e:
        return jsonify({'erro': str(e)}), 500


@app.route('/meus_pedidos/<int:pedido_id>/itens', methods=['GET'])
@app.route('/pedido/<int:pedido_id>/itens', methods=['GET'])
def listar_itens_meu_pedido(pedido_id):
    try:
        if 'user_data' not in session and not request.args.get('usuario_id'):
            return jsonify({'erro': 'Usuário não autenticado'}), 401

        usuario_solicitante = usuario_solicitante_da_requisicao(request.args)
        if not usuario_solicitante:
            return jsonify({'erro': 'Usuário não autenticado'}), 401

        pedido = Pedido.query.filter_by(id=pedido_id).first()
        if not pedido or not pode_acessar_pedido(usuario_solicitante, pedido):
            return jsonify({'erro': 'Pedido não encontrado para este usuário'}), 404

        itens = []
        for item in pedido.itens:
            itens.append({
                'id': item.id,
                'referencia_produto': item.referencia_produto,
                'quantidade': item.quantidade,
                'desconto': round(getattr(item, 'desconto', 0.0) or 0.0, 2),
                'valor_unitario': round(item.valor_unitario, 2),
                'valor_total': round(item.valor_total, 2)
            })

        return jsonify({
            'pedido_id': pedido.id,
            'valor_total': round(pedido.valor_total, 2),
            'itens': itens
        }), 200
    except Exception as e:
        return jsonify({'erro': str(e)}), 500


@app.route('/clientes', methods=['GET'])
def listar_clientes():
    if 'user_data' not in session and not request.args.get('usuario_id'):
        return jsonify({'erro': 'Usuário não autenticado'}), 401

    usuario_solicitante = usuario_solicitante_da_requisicao(request.args)
    if not usuario_solicitante:
        return jsonify({'erro': 'Usuário não autenticado'}), 401

    if usuario_solicitante.is_admin:
        usuarios = Usuario.query.order_by(Usuario.nome.asc()).all()
    else:
        usuarios = [usuario_solicitante]

    return jsonify([serializar_usuario(usuario) for usuario in usuarios]), 200


@app.route('/relatorio_vendas', methods=['GET'])
def relatorio_vendas():
    if 'user_data' not in session and not request.args.get('usuario_id'):
        return jsonify({'erro': 'Usuário não autenticado'}), 401

    usuario_solicitante = usuario_solicitante_da_requisicao(request.args)
    if not usuario_solicitante:
        return jsonify({'erro': 'Usuário não autenticado'}), 401

    query = Pedido.query
    cliente_id = request.args.get('cliente_id', type=int)
    data_inicio_raw = (request.args.get('data_inicio') or '').strip()
    data_fim_raw = (request.args.get('data_fim') or '').strip()

    if usuario_solicitante.is_admin:
        if cliente_id:
            query = query.filter(Pedido.usuario_id == cliente_id)
    else:
        query = query.filter(Pedido.usuario_id == usuario_solicitante.id)

    data_inicio = None
    data_fim = None
    if data_inicio_raw:
        try:
            data_inicio = datetime.strptime(data_inicio_raw, '%Y-%m-%d')
        except ValueError:
            return jsonify({'erro': 'Data inicial inválida. Use YYYY-MM-DD.'}), 400
        query = query.filter(Pedido.data_pedido >= data_inicio)

    if data_fim_raw:
        try:
            data_fim = datetime.strptime(data_fim_raw, '%Y-%m-%d')
        except ValueError:
            return jsonify({'erro': 'Data final inválida. Use YYYY-MM-DD.'}), 400
        data_fim = data_fim.replace(hour=23, minute=59, second=59, microsecond=999999)
        query = query.filter(Pedido.data_pedido <= data_fim)

    pedidos = query.order_by(Pedido.data_pedido.desc()).all()
    total_vendas = round(sum(pedido.valor_total for pedido in pedidos), 2)

    return jsonify({
        'pedidos': [serializar_pedido(pedido) for pedido in pedidos],
        'resumo': {
            'quantidade_pedidos': len(pedidos),
            'valor_total': total_vendas,
            'data_inicio': data_inicio_raw,
            'data_fim': data_fim_raw,
            'cliente_id': cliente_id,
        }
    }), 200


@app.route('/clientes/<int:cliente_id>', methods=['PUT'])
@app.route('/usuarios/<int:cliente_id>', methods=['PUT'])
def atualizar_cliente(cliente_id):
    if 'user_data' not in session and not request.json.get('usuario_id'):
        return jsonify({'erro': 'Usuário não autenticado'}), 401

    payload = request.get_json(silent=True) or request.form
    usuario_solicitante = usuario_solicitante_da_requisicao(payload)
    if not usuario_solicitante:
        return jsonify({'erro': 'Usuário não autenticado'}), 401

    if not usuario_solicitante.is_admin and usuario_solicitante.id != cliente_id:
        return jsonify({'erro': 'Sem permissão para alterar este cadastro'}), 403

    usuario = Usuario.query.filter_by(id=cliente_id).first()
    if not usuario:
        return jsonify({'erro': 'Usuário não encontrado'}), 404

    nome = (payload.get('nome') or usuario.nome).strip()
    email = (payload.get('email') or usuario.email).strip().lower()
    celular = (payload.get('celular') or usuario.celular or '').strip()
    cnpj = normalizar_cnpj(payload.get('cnpj') or usuario.cnpj)
    senha = payload.get('senha')
    empresas = payload.get('empresas', usuario.associacoes)
    empresas_lista = parse_empresas(empresas)
    empresa_ativa = (payload.get('empresa_ativa') or usuario.empresa_ativa or empresas_lista[0]).strip()

    if not all([nome, email, cnpj]):
        return jsonify({'erro': 'Nome, email e CNPJ são obrigatórios.'}), 400

    if len(cnpj) != 14:
        return jsonify({'erro': 'CNPJ inválido. Informe 14 dígitos.'}), 400

    usuario_existente = Usuario.query.filter(
        Usuario.id != usuario.id,
        Usuario.email == email,
        Usuario.cnpj == cnpj,
    ).first()
    if usuario_existente:
        return jsonify({'erro': 'Já existe um usuário com este email e CNPJ.'}), 409

    if celular:
        celular_existente = Usuario.query.filter(
            Usuario.id != usuario.id,
            Usuario.celular == celular,
        ).first()
        if celular_existente:
            return jsonify({'erro': 'Já existe um usuário com este celular.'}), 409

    if empresa_ativa not in empresas_lista:
        empresa_ativa = empresas_lista[0]

    usuario.nome = nome
    usuario.email = email
    usuario.celular = celular or None
    usuario.cnpj = cnpj
    usuario.associacoes = ','.join(empresas_lista)
    usuario.empresa_ativa = empresa_ativa

    if usuario_solicitante.is_admin:
        usuario.perm_download_tabelas = valor_booleano(payload.get('perm_download_tabelas', usuario.perm_download_tabelas), usuario.perm_download_tabelas)
        usuario.perm_fazer_pedido = valor_booleano(payload.get('perm_fazer_pedido', usuario.perm_fazer_pedido), usuario.perm_fazer_pedido)
        usuario.perm_solicitar_visita = valor_booleano(payload.get('perm_solicitar_visita', usuario.perm_solicitar_visita), usuario.perm_solicitar_visita)
        usuario.perm_visualizar_pedidos = valor_booleano(payload.get('perm_visualizar_pedidos', usuario.perm_visualizar_pedidos), usuario.perm_visualizar_pedidos)

    if senha:
        usuario.senha = senha

    db.session.commit()

    usuario_serializado = serializar_usuario(usuario)
    if usuario_solicitante.id == usuario.id:
        session['user_data'] = usuario_serializado

    return jsonify(usuario_serializado), 200


@app.route('/clientes/<int:cliente_id>', methods=['DELETE'])
@app.route('/usuarios/<int:cliente_id>', methods=['DELETE'])
def deletar_cliente(cliente_id):
    if 'user_data' not in session and not request.json.get('usuario_id'):
        return jsonify({'erro': 'Usuário não autenticado'}), 401

    payload = request.get_json(silent=True) or request.args or request.form
    usuario_solicitante = usuario_solicitante_da_requisicao(payload)
    if not usuario_solicitante:
        return jsonify({'erro': 'Usuário não autenticado'}), 401

    if not usuario_solicitante.is_admin:
        return jsonify({'erro': 'Apenas administrador pode deletar clientes.'}), 403

    if usuario_solicitante.id == cliente_id:
        return jsonify({'erro': 'Não é possível deletar o usuário administrador atualmente em uso.'}), 400

    usuario = Usuario.query.filter_by(id=cliente_id).first()
    if not usuario:
        return jsonify({'erro': 'Usuário não encontrado.'}), 404

    try:
        nome_cliente = usuario.nome
        Agendamento.query.filter_by(usuario_id=cliente_id).delete()
        pedidos = Pedido.query.filter_by(usuario_id=cliente_id).all()
        for p in pedidos:
            db.session.delete(p)
        db.session.delete(usuario)
        db.session.commit()
        return jsonify({'mensagem': f'Cliente "{nome_cliente}" deletado com sucesso.'}), 200
    except Exception as e:
        db.session.rollback()
        return jsonify({'erro': f'Erro ao deletar cliente: {str(e)}'}), 500


@app.route('/selecionar_empresa', methods=['POST'])
def selecionar_empresa():
    if 'user_data' not in session and not request.json.get('usuario_id'):
        flash('Faça login para selecionar a empresa.', 'error')
        return redirect('/')

    payload = request.get_json(silent=True) or request.form
    usuario_id = payload.get('usuario_id') or session.get('user_data', {}).get('id')
    empresa = (payload.get('empresa') or payload.get('empresa_ativa') or '').strip()

    if not usuario_id or not empresa:
        if request.is_json:
            return jsonify({'erro': 'usuario_id e empresa são obrigatórios.'}), 400
        flash('Selecione a empresa.', 'error')
        return render_template('index.html', **contexto_index())

    usuario = Usuario.query.filter_by(id=usuario_id).first()
    if not usuario:
        if request.is_json:
            return jsonify({'erro': 'Usuário não encontrado.'}), 404
        flash('Usuário não encontrado.', 'error')
        return render_template('index.html', **contexto_index())

    empresas_usuario = parse_empresas(usuario.associacoes)
    if empresa not in empresas_usuario:
        if request.is_json:
            return jsonify({'erro': 'Usuário não possui acesso à empresa selecionada.'}), 403
        flash('Sem acesso a essa empresa.', 'error')
        return render_template('index.html', **contexto_index())

    usuario.empresa_ativa = empresa
    db.session.commit()

    usuario_serializado = serializar_usuario(usuario)
    session['user_data'] = usuario_serializado

    if request.is_json:
        return jsonify(usuario_serializado), 200

    flash('Empresa ativa alterada com sucesso.', 'success')
    return render_template('index.html', **contexto_index())


if __name__ == '__main__':
    port = int(os.getenv('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False)
